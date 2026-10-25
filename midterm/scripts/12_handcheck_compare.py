"""Stage 13: compare the manual Excel check (results/hand_check_ap.xlsx) with the pipeline code.

User 4949222956972288366, val protocol: history < 2020-05-18, window / GT 2020-05-18 .. 2020-07-26.
GT-D80 and h1 come from scripts/10_h1.py (imported, logic not re-implemented).
No experiments, no existing files are modified; writes only results/handcheck_compare.md.
"""

import importlib.util
import math
import sys
import warnings
from datetime import datetime
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore", category=DeprecationWarning)

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import config  # noqa: E402
import evaluate as E  # noqa: E402
from join import to_unix  # noqa: E402

_spec = importlib.util.spec_from_file_location("h1mod", ROOT / "scripts" / "10_h1.py")
H1 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(H1)

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

USER = 4949222956972288366
REF = to_unix(datetime(2020, 5, 18))
END = to_unix(datetime(2020, 7, 27))
TOL = 1e-4
XLSX = config.RESULTS_DIR / "hand_check_ap.xlsx"
SHOWS_E = [240081, 1100411, 2400467, 2400480, 2400508, 5200303]


def load_excel():
    wv = openpyxl.load_workbook(XLSX, data_only=True)
    wf = openpyxl.load_workbook(XLSX, data_only=False)
    empty = []
    for name in wf.sheetnames:
        for row in wf[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and wv[name][c.coordinate].value is None:
                    empty.append(f"{name}!{c.coordinate}")
    if empty:
        print("STOP: formula cells without cached values (file was not saved in Excel):", ", ".join(empty[:20]),
              f"... total {len(empty)}" if len(empty) > 20 else "")
        sys.exit(1)
    return wv


def main():
    wb = load_excel()
    raw, va, ap, hs = wb["raw_join_example"], wb["val_airings"], wb["AP@5"], wb["h1_score"]
    rows = []  # (block, name, code, excel)

    def add(block, name, code, excel):
        rows.append((block, name, float(code), float(excel)))

    print("Loading pipeline data ...", flush=True)
    matches = E.load_labeled_airings(E.CAP)
    grid = H1.load_grid()
    log = H1.load_log()
    mu = matches[matches["user_id"] == USER]

    # ---- a) join of one airing
    a_st = to_unix(datetime(2020, 5, 18, 18, 45))
    m_a = mu[(mu["tv_show_id"] == 240081) & (mu["channel"] == 5) & (mu["air_start"] == a_st)]
    ov_a = int(m_a["overlap"].sum())
    dur_a = int(m_a["duration"].iloc[0])
    add("a", "overlap, с (240081, канал 5, 2020-05-18 18:45)", ov_a, raw["B9"].value)
    add("a", "pct = overlap / duration", min(ov_a / dur_a, 1.0), raw["D9"].value)
    add("a", "duration эфира, с", dur_a, raw["C3"].value)

    # ---- b) val airings of the user (tv_show_id != 0), overlap summed per airing
    w = mu[(mu["air_start"] >= REF) & (mu["air_start"] < END)]
    g = w.groupby(["tv_show_id", "channel", "air_start"]).agg(ov=("overlap", "sum"), dur=("duration", "first"))
    g["pct"] = (g["ov"] / g["dur"]).clip(0, 1)
    code_pct = {}
    for (sid, _ch, ast), r in g.iterrows():
        code_pct.setdefault((int(sid), int(ast)), []).append(float(r["pct"]))
    xl = []
    for r in range(2, va.max_row + 1):
        sid = va.cell(r, 1).value
        if sid is None:
            continue
        ast = va.cell(r, 4).value
        ast = ast if isinstance(ast, datetime) else datetime.strptime(str(ast), "%Y-%m-%d %H:%M:%S")
        xl.append((int(sid), to_unix(ast), float(va.cell(r, 7).value), int(va.cell(r, 8).value)))
    xl_keys = {(s, a) for s, a, _p, _f in xl}
    diffs, unmatched_x = [], 0
    for s, a, p, _f in xl:
        if (s, a) in code_pct:
            diffs.append(min(abs(p - cp) for cp in code_pct[(s, a)]))
        else:
            unmatched_x += 1
    unmatched_c = sum(len(v) for k, v in code_pct.items() if k not in xl_keys)
    add("b", "число помеченных эфиров в val", len(g), len(xl))
    add("b", "число эфиров с pct >= 0.8", int((g["pct"] >= 0.8).sum()), sum(f for *_x, f in xl))
    add("b", "макс. |Δpct| по строкам (tv_show_id + air_start)", max(diffs) if diffs else float("nan"), 0.0)
    add("b", "строк Excel без пары в коде / строк кода без пары в Excel", unmatched_x + unmatched_c, 0)

    # ---- c) GT-D80 (set from 10_h1.build_gt_d80; counts of pct >= 0.8 airings for order/display)
    gt_set = H1.build_gt_d80(matches, REF, END)[USER]
    cnt = g[g["pct"] >= 0.8].reset_index().groupby("tv_show_id").size()
    gt_order = sorted(gt_set, key=lambda s: (-int(cnt[s]), s))
    for i in range(5):
        add("c", f"GT-D80 ранг {i + 1}: tv_show_id", gt_order[i], ap.cell(2 + i, 11).value)
        add("c", f"GT-D80 ранг {i + 1}: эфиров с pct >= 0.8", int(cnt[gt_order[i]]), ap.cell(2 + i, 12).value)

    # ---- d) h1 top-5 and AP@5 vs D80
    print("h1 rankings (10_h1.h1_rankings) ...", flush=True)
    h1_rank, _gtop = H1.h1_rankings(grid, log, matches, REF, END)
    top5 = h1_rank[USER][:5]
    for i in range(5):
        add("d", f"h1 ранг {i + 1}: tv_show_id", top5[i], ap.cell(2 + i, 2).value)
    add("d", "AP@5 h1 против GT-D80", E.ap_at_k(top5, gt_set), ap["H8"].value)

    # ---- e) p, slot, score, rank for six shows (10_h1 functions: compute_p, slot_profile_from_log, hour_fractions)
    p_all = H1.compute_p(matches, REF)
    prof = H1.slot_profile_from_log(log, REF).get(USER, {})
    sids, chs, sts, durs = grid
    wm = (sts >= REF) & (sts < END)
    slot = {s: 0.0 for s in SHOWS_E}
    for sid, ch, st, du in zip(sids[wm].tolist(), chs[wm].tolist(), sts[wm].tolist(), durs[wm].tolist()):
        if sid in slot:
            for hw, fr in H1.hour_fractions(st, du):
                slot[sid] += prof.get((ch, hw), 0.0) * fr
    score = {s: p_all.get((USER, s), 0.0) * math.sqrt(slot[s] + 1e-4) for s in SHOWS_E}
    order = sorted(SHOWS_E, key=lambda s: (-score[s], s))
    xl_e = {}
    for r in range(2, 8):
        xl_e[int(hs.cell(r, 10).value)] = [hs.cell(r, c).value for c in (11, 12, 13, 14)]
    for s in SHOWS_E:
        k, l_, m, n = xl_e[s]
        add("e", f"{s}: p", p_all.get((USER, s), 0.0), k)
        add("e", f"{s}: slot", slot[s], l_)
        add("e", f"{s}: score", score[s], m)
        add("e", f"{s}: ранг среди шести", order.index(s) + 1, n)

    # ---- report
    md = ["# Этап 13: ручной расчёт (Excel) против кода пайплайна", "",
          f"Пользователь {USER}, протокол val: история < 2020-05-18, окно и GT 2020-05-18 .. 2020-07-26. "
          "GT-D80 и h1 — функции scripts/10_h1.py. Excel: results/hand_check_ap.xlsx, значения формул "
          f"(openpyxl, data_only=True). Допуск {TOL}.", "",
          "Блоки: a — склейка одного эфира (raw_join_example), b — эфиры val (val_airings), "
          "c — эталон D80 (AP@5!K2:L6), d — top-5 h1 и AP@5 (AP@5!B2:B6, H8), e — p / slot / score / ранг (h1_score!K2:N7).",
          "", "Примечание: slot в h1_score!L2:L7 — числа, вставленные из скрипта, а не формулы Excel; "
          "для slot сверка проверяет воспроизводимость, а не независимый ручной расчёт.", "",
          "| блок | величина | код | Excel | \\|разница\\| | совпадает (допуск 1e-4) |", "|---|---|---|---|---|---|"]
    n_ok = 0
    bad = []
    for blk, name, c, x in rows:
        d = abs(c - x)
        ok = d <= TOL
        n_ok += ok
        if not ok:
            bad.append((blk, name, c, x, d))
        fmt = (lambda v: f"{v:.0f}") if float(c).is_integer() and float(x).is_integer() else (lambda v: f"{v:.6f}")
        md.append(f"| {blk} | {name.replace('|', chr(92) + '|')} | {fmt(c)} | {fmt(x)} | {d:.2e} | {'да' if ok else '**нет**'} |")
    md += ["", f"Итог: сверено {len(rows)} величин, совпало {n_ok}, расхождений {len(rows) - n_ok}."]
    if bad:
        md += ["", "## Расхождения (не исправлялись)", ""]
        md += [f"- {b}: {nm}: код {c}, Excel {x}, |разница| {d:.3g}" for b, nm, c, x, d in bad]
    (config.RESULTS_DIR / "handcheck_compare.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print("\n".join(md))
    sys.exit(0 if not bad else 2)


if __name__ == "__main__":
    main()
